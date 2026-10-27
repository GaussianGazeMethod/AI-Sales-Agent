"""闭环三 · 运营服务（P2-b）：运营记录（回访/培训/需求登记）+ 运营指标（活跃率）。

指标口径（spec/04 §3）：
- **活跃率 = 登陆人数 ÷ 总账号数**，需要**日活 / 周活 / 月活**三种
- 取数方式：**先手工从平台后台导出、导入系统**（不打通产品埋点）
- 「高质量文件」由**人工标记**
"""

from __future__ import annotations

import io
import logging
import re
from datetime import date, datetime
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import dicts
from app.core.errors import NotFoundError, ValidationError
from app.models.delivery import Contract
from app.models.enums import MetricPeriodType, OperationActionType, RecordStatus
from app.models.lead import Lead
from app.models.operation import OperationMetric, OperationRecord
from app.models.user import User
from app.schemas import (
    OperationMetricCreate,
    OperationMetricUpdate,
    OperationRecordCreate,
    OperationRecordUpdate,
)
from app.services import audit_service

logger = logging.getLogger(__name__)

METRIC_TEMPLATE_HEADERS = [
    "客户名称",
    "合同编号",
    "统计周期",
    "周期值",
    "总账号数",
    "登陆人数",
    "上传文件次数",
    "下载文件次数",
    "文件总数",
    "高质量文件数量",
]

METRIC_SAMPLE_ROW = [
    "示例：XX市第三中学",
    "",
    "月活",
    "2026-09",
    1200,
    480,
    320,
    210,
    1500,
    36,
]

_PERIOD_TYPE_ALIAS = {
    "日活": MetricPeriodType.DAY,
    "日": MetricPeriodType.DAY,
    "day": MetricPeriodType.DAY,
    "周活": MetricPeriodType.WEEK,
    "周": MetricPeriodType.WEEK,
    "week": MetricPeriodType.WEEK,
    "月活": MetricPeriodType.MONTH,
    "月": MetricPeriodType.MONTH,
    "month": MetricPeriodType.MONTH,
}

_PERIOD_RE = {
    MetricPeriodType.DAY: re.compile(r"^\d{4}-\d{2}-\d{2}$"),
    MetricPeriodType.WEEK: re.compile(r"^\d{4}-W\d{2}$"),
    MetricPeriodType.MONTH: re.compile(r"^\d{4}-\d{2}$"),
}


def calc_activity_rate(login_count: int, total_accounts: int) -> float:
    """活跃率 = 登陆人数 ÷ 总账号数。"""
    if not total_accounts:
        return 0.0
    return round(login_count / total_accounts, 4)


def validate_period(period_type: str, period: str) -> None:
    pattern = _PERIOD_RE.get(period_type)
    if pattern is None:
        raise ValidationError("统计周期只能是 day / week / month")
    if not pattern.match(period or ""):
        example = {
            MetricPeriodType.DAY: "2026-09-30",
            MetricPeriodType.WEEK: "2026-W40",
            MetricPeriodType.MONTH: "2026-09",
        }[period_type]
        raise ValidationError(f"周期值格式应为 {example}")


# ══════════════ 运营记录（3.6 需求登记 / 3.10 回访）══════════════


async def list_records(
    db: AsyncSession,
    *,
    lead_id: str | None = None,
    contract_id: str | None = None,
    action_type: str | None = None,
    status: str | None = None,
    start: date | None = None,
    end: date | None = None,
    offset: int = 0,
    limit: int = 20,
) -> list[OperationRecord]:
    stmt = select(OperationRecord)
    if lead_id:
        stmt = stmt.where(OperationRecord.lead_id == lead_id)
    if contract_id:
        stmt = stmt.where(OperationRecord.contract_id == contract_id)
    if action_type:
        stmt = stmt.where(OperationRecord.action_type == action_type)
    if status:
        stmt = stmt.where(OperationRecord.status == status)
    if start:
        stmt = stmt.where(OperationRecord.record_date >= start)
    if end:
        stmt = stmt.where(OperationRecord.record_date <= end)
    return list(
        (
            await db.execute(
                stmt.order_by(OperationRecord.record_date.desc(), OperationRecord.created_at.desc())
                .offset(offset)
                .limit(limit)
            )
        )
        .scalars()
        .all()
    )


async def count_records(db: AsyncSession, **filters) -> int:
    stmt = select(func.count()).select_from(OperationRecord)
    if filters.get("contract_id"):
        stmt = stmt.where(OperationRecord.contract_id == filters["contract_id"])
    if filters.get("lead_id"):
        stmt = stmt.where(OperationRecord.lead_id == filters["lead_id"])
    if filters.get("action_type"):
        stmt = stmt.where(OperationRecord.action_type == filters["action_type"])
    if filters.get("status"):
        stmt = stmt.where(OperationRecord.status == filters["status"])
    return int((await db.execute(stmt)).scalar() or 0)


async def add_record(
    db: AsyncSession, *, operator: User, payload: OperationRecordCreate
) -> OperationRecord:
    lead = (
        await db.execute(select(Lead).where(Lead.id == payload.lead_id, Lead.is_deleted.is_(False)))
    ).scalar_one_or_none()
    if lead is None:
        raise NotFoundError("客户不存在")

    contract_id = payload.contract_id
    if not contract_id:
        from app.services import delivery_service

        contract = await delivery_service.get_contract_by_lead(db, lead.id)
        contract_id = contract.id if contract else None

    record = OperationRecord(
        lead_id=lead.id,
        contract_id=contract_id,
        customer_name=lead.name,
        record_date=payload.record_date or date.today(),
        action_type=payload.action_type,
        content=payload.content,
        next_plan=payload.next_plan,
        status=payload.status,
        satisfaction=payload.satisfaction,
        executor_id=operator.id,
        executor_name=operator.name,
        source="manual",
    )
    db.add(record)
    await db.flush()
    await audit_service.write_log(
        db,
        operator=operator,
        action="operation_record",
        target_type="lead",
        target_id=lead.id,
        target_name=lead.name,
        detail=f"运营记录（{dicts.label_of(dicts.OPERATION_ACTION_LABELS, payload.action_type)}）："
        f"{payload.content[:60]}",
    )
    return record


async def update_record(
    db: AsyncSession, *, operator: User, record: OperationRecord, payload: OperationRecordUpdate
) -> OperationRecord:
    for key, value in payload.model_dump(exclude_unset=True).items():
        if value is not None:
            setattr(record, key, value)
    return record


async def get_record(db: AsyncSession, record_id: str) -> OperationRecord:
    row = (
        await db.execute(select(OperationRecord).where(OperationRecord.id == record_id))
    ).scalar_one_or_none()
    if row is None:
        raise NotFoundError("运营记录不存在")
    return row


async def delete_record(db: AsyncSession, *, operator: User, record: OperationRecord) -> None:
    await audit_service.write_log(
        db,
        operator=operator,
        action="delete",
        target_type="operation_record",
        target_id=record.id,
        target_name=record.customer_name,
        detail="删除运营记录",
    )
    await db.delete(record)


async def handover_confirmed_at(db: AsyncSession, contract: Contract) -> datetime | None:
    """双方交接确认的实际时点；缺确认记录或时点时不启动运营计时。"""
    from app.models.delivery import HandoverSheet

    handover = (await db.execute(select(HandoverSheet).where(
        HandoverSheet.contract_id == contract.id
    ))).scalar_one_or_none()
    if (handover is None or not handover.is_confirmed
            or handover.sales_confirmed_at is None
            or handover.delivery_confirmed_at is None):
        return None
    return max(handover.sales_confirmed_at, handover.delivery_confirmed_at)


async def last_contact_at(db: AsyncSession, contract: Contract) -> datetime | None:
    """只认本合同交接后已完成的运营联系，不用销售跟进或其他合同重置计时。"""
    confirmed = await handover_confirmed_at(db, contract)
    if confirmed is None:
        return None
    last_record = (await db.execute(select(func.max(OperationRecord.record_date)).where(
        OperationRecord.contract_id == contract.id,
        OperationRecord.record_date >= confirmed.date(),
        OperationRecord.status == "done",
        OperationRecord.content != "",
    ))).scalar()
    return datetime.combine(last_record, datetime.min.time()) if last_record else None


async def contact_rhythm(db, contract, *, today=None, threshold_days=None):
    """合同页面和预警扫描共用运营联系时钟。"""
    from app.services import risk_service

    confirmed = await handover_confirmed_at(db, contract)
    contact = await last_contact_at(db, contract)
    if threshold_days is None:
        threshold_days = int((await risk_service.get_rules(db))["ops_no_contact_days"])
    since = contact or confirmed
    days = ((today or date.today()) - since.date()).days if since else None
    return {
        "confirmed_at": confirmed.isoformat() if confirmed else None,
        "last_contact_at": contact.isoformat() if contact else None,
        "idle_days": days,
        "threshold_days": threshold_days,
        "overdue": days is not None and days >= threshold_days,
    }


def record_out(record: OperationRecord) -> dict:
    return {
        "id": record.id,
        "lead_id": record.lead_id,
        "contract_id": record.contract_id,
        "customer_name": record.customer_name,
        "record_date": record.record_date.isoformat() if record.record_date else None,
        "action_type": record.action_type,
        "action_type_label": dicts.label_of(dicts.OPERATION_ACTION_LABELS, record.action_type),
        "content": record.content,
        "next_plan": record.next_plan,
        "status": record.status,
        "status_label": dicts.label_of(dicts.RECORD_STATUS_LABELS, record.status),
        "satisfaction": record.satisfaction,
        "executor_id": record.executor_id,
        "executor_name": record.executor_name,
        "source": record.source,
        "batch_no": record.batch_no,
        "created_at": record.created_at.isoformat() if record.created_at else None,
    }


# ══════════════ 运营指标（3.9）══════════════


async def list_metrics(
    db: AsyncSession,
    *,
    lead_id: str | None = None,
    contract_id: str | None = None,
    period_type: str | None = None,
    period: str | None = None,
    offset: int = 0,
    limit: int = 50,
) -> list[OperationMetric]:
    stmt = select(OperationMetric)
    if lead_id:
        stmt = stmt.where(OperationMetric.lead_id == lead_id)
    if contract_id:
        stmt = stmt.where(OperationMetric.contract_id == contract_id)
    if period_type:
        stmt = stmt.where(OperationMetric.period_type == period_type)
    if period:
        stmt = stmt.where(OperationMetric.period == period)
    return list(
        (await db.execute(stmt.order_by(OperationMetric.period.desc()).offset(offset).limit(limit)))
        .scalars()
        .all()
    )


async def count_metrics(db: AsyncSession, **filters) -> int:
    stmt = select(func.count()).select_from(OperationMetric)
    for key in ("lead_id", "contract_id", "period_type", "period"):
        if filters.get(key):
            stmt = stmt.where(getattr(OperationMetric, key) == filters[key])
    return int((await db.execute(stmt)).scalar() or 0)


async def add_metric(
    db: AsyncSession, *, operator: User, payload: OperationMetricCreate
) -> OperationMetric:
    lead = (
        await db.execute(select(Lead).where(Lead.id == payload.lead_id, Lead.is_deleted.is_(False)))
    ).scalar_one_or_none()
    if lead is None:
        raise NotFoundError("客户不存在")
    validate_period(payload.period_type, payload.period)

    contract_id = payload.contract_id
    if not contract_id:
        from app.services import delivery_service

        contract = await delivery_service.get_contract_by_lead(db, lead.id)
        contract_id = contract.id if contract else None

    row = OperationMetric(
        lead_id=lead.id,
        contract_id=contract_id,
        customer_name=lead.name,
        period_type=payload.period_type,
        period=payload.period,
        total_accounts=payload.total_accounts,
        login_count=payload.login_count,
        activity_rate=calc_activity_rate(payload.login_count, payload.total_accounts),
        upload_count=payload.upload_count,
        download_count=payload.download_count,
        file_count=payload.file_count,
        high_quality_file_count=payload.high_quality_file_count,
        source="manual",
        imported_by_id=operator.id,
        imported_by_name=operator.name,
    )
    db.add(row)
    await db.flush()
    return row


async def update_metric(
    db: AsyncSession, *, operator: User, row: OperationMetric, payload: OperationMetricUpdate
) -> OperationMetric:
    data = payload.model_dump(exclude_unset=True)
    for key, value in data.items():
        if value is not None:
            setattr(row, key, value)
    row.activity_rate = calc_activity_rate(row.login_count, row.total_accounts)
    return row


async def get_metric(db: AsyncSession, metric_id: str) -> OperationMetric:
    row = (
        await db.execute(select(OperationMetric).where(OperationMetric.id == metric_id))
    ).scalar_one_or_none()
    if row is None:
        raise NotFoundError("运营指标不存在")
    return row


def build_metric_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "运营指标导入"
    ws.append(METRIC_TEMPLATE_HEADERS)
    ws.append(METRIC_SAMPLE_ROW)
    for idx, header in enumerate(METRIC_TEMPLATE_HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(
            12, len(header) * 2 + 4
        )

    tips = wb.create_sheet("填写说明")
    tips.append(["列名", "是否必填", "说明"])
    tips.append(["客户名称", "必填", "与系统中的客户（线索/商机）名称一致，否则该行失败"])
    tips.append(["统计周期", "必填", "日活 / 周活 / 月活（也可写 day/week/month）"])
    tips.append(["周期值", "必填", "日活：2026-09-30；周活：2026-W40；月活：2026-09"])
    tips.append(["总账号数/登陆人数", "必填", "活跃率 = 登陆人数 ÷ 总账号数，由系统自动计算"])
    tips.append(["高质量文件数量", "选填", "人工标记的「高质量文件」数量"])
    tips.append(["合同编号", "选填", "填了则挂到该合同下，不填则挂到该客户最新合同"])
    tips.column_dimensions["A"].width = 22
    tips.column_dimensions["B"].width = 12
    tips.column_dimensions["C"].width = 80

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _ws(value) -> str:
    return "" if value is None else str(value).strip()


async def import_metrics(
    db: AsyncSession, *, operator: User, filename: str, content: bytes
) -> dict:
    """导入运营指标（Excel）：逐行校验 → 入库 → 失败明细。"""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValidationError(f"无法解析 Excel 文件：{exc}") from exc

    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValidationError("Excel 内容为空")

    header = [_ws(c) for c in rows[0]]
    missing = {"客户名称", "统计周期", "周期值"} - set(header)
    if missing:
        raise ValidationError("表头缺少必需列：" + "、".join(sorted(missing)))

    idx = {name: i for i, name in enumerate(header)}
    batch_no = f"OPM{datetime.now().strftime('%Y%m%d%H%M%S')}"

    def cell(row, name: str) -> str:
        pos = idx.get(name)
        if pos is None or pos >= len(row):
            return ""
        return _ws(row[pos])

    total = success = 0
    errors: list[dict] = []
    for row_no, row in enumerate(rows[1:], start=2):
        if row is None or all(_ws(c) == "" for c in row):
            continue
        total += 1
        name = cell(row, "客户名称")
        if name.startswith("示例："):
            total -= 1
            continue
        if not name:
            errors.append({"row": row_no, "message": "客户名称为空"})
            continue

        period_type = _PERIOD_TYPE_ALIAS.get(
            cell(row, "统计周期").lower()
        ) or _PERIOD_TYPE_ALIAS.get(cell(row, "统计周期"))
        if not period_type:
            errors.append(
                {
                    "row": row_no,
                    "message": f"统计周期无法识别：{cell(row, '统计周期')}（日活/周活/月活）",
                }
            )
            continue
        period = cell(row, "周期值")

        lead = (
            (await db.execute(select(Lead).where(Lead.name == name, Lead.is_deleted.is_(False))))
            .scalars()
            .first()
        )
        if lead is None:
            errors.append({"row": row_no, "message": f"客户不存在：{name}（请先录入线索/商机）"})
            continue

        contract = None
        contract_no = cell(row, "合同编号")
        if contract_no:
            contract = (
                await db.execute(select(Contract).where(Contract.contract_no == contract_no))
            ).scalar_one_or_none()
            if contract is None:
                errors.append({"row": row_no, "message": f"合同编号不存在：{contract_no}"})
                continue
        else:
            from app.services import delivery_service

            contract = await delivery_service.get_contract_by_lead(db, lead.id)

        def num(field: str) -> int:
            raw = cell(row, field)
            if not raw:
                return 0
            try:
                return int(float(raw))
            except ValueError:
                return 0

        try:
            validate_period(period_type, period)
        except ValidationError as exc:
            errors.append({"row": row_no, "message": exc.message})
            continue

        total_accounts = num("总账号数")
        login_count = num("登陆人数")
        db.add(
            OperationMetric(
                lead_id=lead.id,
                contract_id=contract.id if contract else None,
                customer_name=lead.name,
                period_type=period_type,
                period=period,
                total_accounts=total_accounts,
                login_count=login_count,
                activity_rate=calc_activity_rate(login_count, total_accounts),
                upload_count=num("上传文件次数"),
                download_count=num("下载文件次数"),
                file_count=num("文件总数"),
                high_quality_file_count=num("高质量文件数量"),
                source="excel_import",
                batch_no=batch_no,
                imported_by_id=operator.id,
                imported_by_name=operator.name,
            )
        )
        success += 1

    await db.flush()
    await audit_service.write_log(
        db,
        operator=operator,
        action="import",
        target_type="operation_metric",
        target_name=filename,
        detail=(
            f"运营指标导入：共 {total} 条，成功 {success}，失败 {total - success}"
            f"（批次 {batch_no}）"
        ),
    )
    return {
        "batch_no": batch_no,
        "total": total,
        "success": success,
        "failed": total - success,
        "errors": errors[:200],
    }


async def metric_dashboard(
    db: AsyncSession,
    *,
    period_type: str = MetricPeriodType.MONTH,
    period: str | None = None,
    owner_id: str | None = None,
) -> dict:
    """活跃率看板：总览 + 分学校排名 + 近 6 期趋势。"""
    latest_period = (
        period
        or (
            await db.execute(
                select(func.max(OperationMetric.period)).where(
                    OperationMetric.period_type == period_type
                )
            )
        ).scalar()
    )

    stmt = select(OperationMetric).where(OperationMetric.period_type == period_type)
    if latest_period:
        stmt = stmt.where(OperationMetric.period == latest_period)
    if owner_id:
        stmt = stmt.where(
            OperationMetric.contract_id.in_(
                select(Contract.id).where(Contract.delivery_owner_id == owner_id)
            )
        )
    rows = list((await db.execute(stmt)).scalars().all())

    total_accounts = sum(r.total_accounts for r in rows)
    total_login = sum(r.login_count for r in rows)
    high_quality = sum(r.high_quality_file_count for r in rows)
    ranking = sorted(rows, key=lambda r: r.activity_rate, reverse=True)
    ranking_out = [
        {
            **metric_out(r),
            "rank": idx + 1,
        }
        for idx, r in enumerate(ranking)
    ]

    # 趋势（近 6 期）
    trend_rows = (
        await db.execute(
            select(
                OperationMetric.period,
                func.avg(OperationMetric.activity_rate),
                func.sum(OperationMetric.login_count),
                func.count(),
            )
            .where(OperationMetric.period_type == period_type)
            .group_by(OperationMetric.period)
            .order_by(OperationMetric.period.desc())
            .limit(6)
        )
    ).all()

    return {
        "period_type": period_type,
        "period_type_label": dicts.label_of(dicts.METRIC_PERIOD_LABELS, period_type),
        "period": latest_period,
        "overview": {
            "customer_count": len(rows),
            "total_accounts": total_accounts,
            "total_login": total_login,
            "avg_activity_rate": calc_activity_rate(total_login, total_accounts),
            "high_quality_file_count": high_quality,
        },
        "ranking": ranking_out,
        "trend": [
            {
                "period": r[0],
                "avg_activity_rate": round(float(r[1] or 0), 4),
                "login_count": int(r[2] or 0),
                "customer_count": int(r[3] or 0),
            }
            for r in sorted(trend_rows, key=lambda x: x[0])
        ],
    }


def metric_out(row: OperationMetric) -> dict:
    return {
        "id": row.id,
        "lead_id": row.lead_id,
        "contract_id": row.contract_id,
        "customer_name": row.customer_name,
        "period_type": row.period_type,
        "period_type_label": dicts.label_of(dicts.METRIC_PERIOD_LABELS, row.period_type),
        "period": row.period,
        "total_accounts": row.total_accounts,
        "login_count": row.login_count,
        "activity_rate": row.activity_rate,
        "upload_count": row.upload_count,
        "download_count": row.download_count,
        "file_count": row.file_count,
        "high_quality_file_count": row.high_quality_file_count,
        "source": row.source,
        "batch_no": row.batch_no,
        "imported_by_name": row.imported_by_name,
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


async def low_activity_customers(
    db: AsyncSession,
    *,
    period_type: str = MetricPeriodType.MONTH,
    threshold: float = 0.2,
    limit: int = 20,
) -> list[dict]:
    """低活跃客户清单（活跃率不设阈值预警，仅用于复盘展示）。"""
    rows = await list_metrics(db, period_type=period_type, limit=500)
    low = [r for r in rows if r.activity_rate < threshold]
    low.sort(key=lambda r: r.activity_rate)
    return [metric_out(r) for r in low[:limit]]


# ══════════════ 运营记录批量导入（历史数据迁移）══════════════
#
# 表头口径：对标飞书「运营交付部」在线表格的回访/培训记录列。

RECORD_HEADERS = [
    "客户名称",
    "合同编号",
    "记录日期",
    "动作类型",
    "记录内容",
    "下次计划",
    "处理状态",
    "满意度",
    "执行人",
]

RECORD_SAMPLE_ROW = [
    "示例：XX市第三中学",
    "",
    "2026-09-18",
    "定期回访",
    "月度回访一线老师，反馈 AI 阅卷识别准确率良好；希望增加学科模块",
    "下月继续跟进学科模块需求",
    "已完成",
    "满意",
    "演示交付人员",
]

RECORD_ALIASES: dict[str, list[str]] = {
    "customer_name": ["客户名称", "客户", "学校名称"],
    "contract_no": ["合同编号", "合同号"],
    "record_date": ["记录日期", "日期", "回访日期", "发生日期"],
    "action_type": ["动作类型", "类型", "服务类型", "动作"],
    "content": ["记录内容", "内容", "回访内容", "服务内容", "说明"],
    "next_plan": ["下次计划", "下一步", "后续计划"],
    "status": ["处理状态", "状态"],
    "satisfaction": ["满意度", "评价"],
    "executor_name": ["执行人", "负责人", "服务人"],
}

#: 动作类型别名 → 内部代码
_ACTION_ALIAS: dict[str, str] = {}
for _code, _label in dicts.OPERATION_ACTION_LABELS.items():
    _ACTION_ALIAS[_label] = _code
    _ACTION_ALIAS[_code] = _code
_ACTION_ALIAS.update(
    {
        "回访": OperationActionType.VISIT,
        "定期回访": OperationActionType.VISIT,
        "培训": OperationActionType.TRAINING,
        "使用培训": OperationActionType.TRAINING,
        "数据报告": OperationActionType.REPORT,
        "运营报告": OperationActionType.REPORT,
        "活动": OperationActionType.ACTIVITY,
        "活动运营": OperationActionType.ACTIVITY,
        "需求": OperationActionType.REQUIREMENT,
        "需求收集": OperationActionType.REQUIREMENT,
        "需求登记": OperationActionType.REQUIREMENT,
        "投诉跟进": OperationActionType.COMPLAINT_FOLLOW,
        "投诉处理": OperationActionType.COMPLAINT_FOLLOW,
    }
)

#: 处理状态别名 → 内部代码
_STATUS_ALIAS: dict[str, str] = {
    "待处理": RecordStatus.PENDING,
    "未处理": RecordStatus.PENDING,
    "pending": RecordStatus.PENDING,
    "已完成": RecordStatus.DONE,
    "完成": RecordStatus.DONE,
    "已处理": RecordStatus.DONE,
    "done": RecordStatus.DONE,
}


def build_record_template() -> bytes:
    """运营记录导入模板（含填写说明）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "运营记录导入"
    ws.append(RECORD_HEADERS)
    ws.append(RECORD_SAMPLE_ROW)
    for idx, header in enumerate(RECORD_HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(
            12, len(header) * 2 + 4
        )

    tips = wb.create_sheet("填写说明")
    tips.append(["列名", "是否必填", "说明"])
    tips.append(["客户名称", "必填", "与系统中的客户名称一致；填了合同编号时优先按编号匹配"])
    tips.append(["合同编号", "选填", "精确匹配；客户名称重复时请用它"])
    tips.append(["记录日期", "必填", "YYYY-MM-DD（也认 2026/9/18、2026年9月18日）"])
    tips.append(
        [
            "动作类型",
            "必填",
            "定期回访 / 使用培训 / 数据报告 / 活动运营 / 需求收集 / 投诉跟进"
            "（也认 visit/training/report/activity/requirement/complaint_follow）",
        ]
    )
    tips.append(["记录内容", "必填", "写清可核对的事实（谁、做了什么、约定了什么）"])
    tips.append(["下次计划", "选填", ""])
    tips.append(["处理状态", "选填", "待处理 / 已完成（默认已完成）"])
    tips.append(["满意度", "选填", "文本，如 满意 / 一般 / 不满意"])
    tips.append(["执行人", "选填", "填系统用户姓名则记为该执行人；否则按导入人记录"])
    tips.column_dimensions["A"].width = 22
    tips.column_dimensions["B"].width = 14
    tips.column_dimensions["C"].width = 90

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


async def import_records(
    db: AsyncSession, *, operator: User, filename: str, content: bytes
) -> dict:
    """导入运营记录（Excel）：逐行校验 → 入库 → 失败明细。"""
    from app.services import import_service

    header, rows = import_service.read_first_sheet(content)
    _, index = import_service.cell_reader(header, RECORD_ALIASES, ())
    if "content" not in index:
        raise ValidationError("表头缺少必需列：记录内容")
    if "action_type" not in index:
        raise ValidationError("表头缺少必需列：动作类型")
    if "customer_name" not in index and "contract_no" not in index:
        raise ValidationError("表头缺少必需列：客户名称 或 合同编号")

    batch_no = f"OPR{datetime.now().strftime('%Y%m%d%H%M%S')}{uuid4().hex[:4].upper()}"
    total = success = 0
    errors: list[dict] = []

    for row_no, row in enumerate(rows, start=2):
        if row is None or all(import_service._cell_text(c) == "" for c in row):
            continue
        total += 1
        cell, _ = import_service.cell_reader(header, RECORD_ALIASES, row)

        customer_name = cell("customer_name")
        if customer_name.startswith("示例："):
            total -= 1
            continue
        content_text = cell("content")
        if not content_text:
            errors.append({"row": row_no, "message": "记录内容为空"})
            continue

        action_raw = cell("action_type")
        action_type = _ACTION_ALIAS.get(action_raw) or _ACTION_ALIAS.get(action_raw.lower())
        if not action_type:
            errors.append(
                {
                    "row": row_no,
                    "message": (
                        f"动作类型无法识别：{action_raw}"
                        "（定期回访/使用培训/数据报告/活动运营/需求收集/投诉跟进）"
                    ),
                }
            )
            continue

        date_raw = cell("record_date")
        record_date = import_service.parse_date_cell(date_raw)
        if date_raw and record_date is None:
            errors.append(
                {"row": row_no, "message": f"记录日期无法识别：{date_raw}（应为 YYYY-MM-DD）"}
            )
            continue
        if record_date is None:
            errors.append({"row": row_no, "message": "记录日期为空"})
            continue

        lead, row_contract, error = await import_service.resolve_customer(
            db, contract_no=cell("contract_no"), customer_name=customer_name
        )
        if error:
            errors.append({"row": row_no, "message": error})
            continue

        status_raw = cell("status")
        status = _STATUS_ALIAS.get(status_raw) or _STATUS_ALIAS.get(status_raw.lower())
        if not status:
            status = RecordStatus.DONE  # 历史记录默认已完成

        executor_name = cell("executor_name")
        executor = await import_service._user_by_name(db, executor_name) if executor_name else None

        db.add(
            OperationRecord(
                lead_id=lead.id,
                contract_id=row_contract.id if row_contract else None,
                customer_name=lead.name,
                record_date=record_date,
                action_type=action_type,
                content=content_text,
                next_plan=cell("next_plan") or None,
                status=status,
                satisfaction=cell("satisfaction") or None,
                executor_id=executor.id if executor else operator.id,
                executor_name=executor.name if executor else (executor_name or operator.name),
                source="excel_import",
                batch_no=batch_no,
            )
        )
        success += 1

    await db.flush()
    await audit_service.write_log(
        db,
        operator=operator,
        action="import",
        target_type="operation_record",
        target_name=filename,
        detail=(
            f"运营记录导入：共 {total} 条，成功 {success}，失败 {total - success}"
            f"（批次 {batch_no}）"
        ),
    )
    return {
        "batch_no": batch_no,
        "total": total,
        "success": success,
        "failed": total - success,
        "errors": errors[:200],
    }
